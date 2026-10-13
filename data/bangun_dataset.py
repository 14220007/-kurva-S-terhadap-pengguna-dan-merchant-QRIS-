"""Bangun Dataset_QRIS_Indonesia_2021-2026.xlsx dan CSV turunannya.

Satu sumber kebenaran: semua angka dan sumbernya ditulis di sini, lalu diekspor ke
Excel (dengan status verifikasi dan tautan) dan ke CSV yang dipakai skrip model.
Status sumber:
  RESMI BI    = angka tertulis langsung di siaran pers Bank Indonesia (bi.go.id)
  PEJABAT BI  = pernyataan pejabat BI (konferensi pers/acara), dikutip media
  ASPI        = data Asosiasi Sistem Pembayaran Indonesia, dikutip media
  TURUNAN     = dihitung dari angka resmi
"""
import csv, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else "."
BI = "https://www.bi.go.id/id/publikasi/ruang-media/news-release/Pages/{}.aspx"
BIEN = "https://www.bi.go.id/en/publikasi/ruang-media/news-release/Pages/{}.aspx"

# periode, pengguna (juta), merchant (juta), status, no. rilis, tanggal rilis, sumber, url, catatan
PM = [
    ("2021-12", 12.80, None, "TURUNAN", "", "", "BI via Bisnis.com 24/02/2023", "https://finansial.bisnis.com/read/20230224/90/1631430/cuan-bri-bbri-catatkan-nilai-transaksi-qris-rp74-triliun-sepanjang-2022", "28,75 juta (Des-2022) dikurangi tambahan 15,95 juta selama 2022"),
    ("2022-02", 15.99, None, "PEJABAT BI", "", "2022-04-04", "Kepala DKSP BI, paparan virtual, via Fortune IDN", "https://www.fortuneidn.com/finance/transaksi-digital-meningkat-bi-bidik-pengguna-qris-capai-26-juta-00-ccw2k-hclm40", ""),
    ("2022-09", 25.20, 21.60, "PEJABAT BI", "", "2022-10-28", "Asisten Gubernur BI via Kontan", "https://keuangan.kontan.co.id/news/makin-diminati-transaksi-qris-capai-rp-297-triliun-hingga-september-2022", ""),
    ("2022-12", 28.75, 22.70, "PEJABAT BI", "", "2023-02", "BI via Bisnis.com 24/02/2023", "https://finansial.bisnis.com/read/20230224/90/1631430/cuan-bri-bbri-catatkan-nilai-transaksi-qris-rp74-triliun-sepanjang-2022", "ASPI mencatat 23,97 juta merchant Des-2022 (Gusti, 2024)"),
    ("2023-02", 30.87, 24.90, "RESMI BI", "25/93/DKom", "2023-04-11", "Siaran Pers BI", BI.format("sp_259323"), ""),
    ("2023-05", 35.80, 26.10, "RESMI BI", "25/165/DKom", "2023-06-22", "Siaran Pers RDG Juni 2023", BIEN.format("sp_2516523"), "Posisi tersirat Mei 2023"),
    ("2023-06", None, 26.70, "RESMI BI", "25/211/DKom", "2023-08-03", "Siaran Pers BI (penghargaan QRIS)", BI.format("sp_2521123"), "91,4% merchant UMKM"),
    ("2023-07", 38.24, 27.51, "RESMI BI", "25/233/DKom", "2023-08-24", "Siaran Pers RDG Agustus 2023", BI.format("sp_2523323"), "Posisi tersirat Juli 2023"),
    ("2023-08", 40.05, 28.38, "RESMI BI", "25/259/DKom", "2023-09-21", "Siaran Pers RDG September 2023", BIEN.format("sp_2525923"), "Posisi tersirat Agustus 2023"),
    ("2023-09", 41.84, 29.04, "RESMI BI", "25/283/DKom", "2023-10-19", "Siaran Pers RDG Oktober 2023", BI.format("sp_2528323"), "Posisi triwulan III 2023"),
    ("2023-10", 43.44, 29.63, "RESMI BI", "25/317/DKom", "2023-11-23", "Siaran Pers RDG November 2023", BIEN.format("sp_2531723"), "Posisi tersirat Oktober 2023"),
    ("2023-11", 45.03, 30.12, "RESMI BI", "25/342/DKom", "2023-12-21", "Siaran Pers RDG Desember 2023", BI.format("sp_2534223"), ""),
    ("2023-12", 45.78, 30.41, "RESMI BI", "26/9/DKom", "2024-01-17", "Siaran Pers RDG Januari 2024", BI.format("sp_260924"), "Kinerja 2023"),
    ("2024-01", 46.37, 30.88, "RESMI BI", "26/33/DKom", "2024-02-21", "Siaran Pers RDG Februari 2024", BI.format("sp_263324"), ""),
    ("2024-02", 46.98, 31.27, "RESMI BI", "26/54/DKom", "2024-03-20", "Siaran Pers RDG Maret 2024", BI.format("sp_265424"), "Posisi tersirat Februari 2024"),
    ("2024-03", 48.12, 31.61, "RESMI BI", "26/80/DKom", "2024-04-24", "Siaran Pers RDG April 2024", BI.format("sp_268024"), "Posisi triwulan I 2024"),
    ("2024-04", 48.90, 31.86, "RESMI BI", "26/108/DKom", "2024-05-22", "Siaran Pers RDG Mei 2024", BI.format("sp_2610824"), ""),
    ("2024-05", 49.76, 32.25, "RESMI BI", "26/125/DKom", "2024-06-20", "Siaran Pers RDG Juni 2024", BIEN.format("sp_2612524"), ""),
    ("2024-06", 50.50, 32.71, "RESMI BI", "26/144/DKom", "2024-07-17", "Siaran Pers RDG Juli 2024", BIEN.format("sp_2614424"), "Posisi triwulan II 2024"),
    ("2024-08", 52.55, 33.77, "RESMI BI", "26/201/DKom", "2024-09-18", "Siaran Pers RDG September 2024", BI.format("sp_2620124"), ""),
    ("2024-09", 53.30, 34.23, "RESMI BI", "26/225/DKom", "2024-10-16", "Siaran Pers RDG Oktober 2024", BI.format("sp_2622524"), "Posisi triwulan III 2024 (versi sebelumnya memakai 53,0 dari media)"),
    ("2024-10", 54.10, 34.70, "RESMI BI", "26/254/DKom", "2024-11-20", "Siaran Pers RDG November 2024", BI.format("sp_2625424"), ""),
    ("2024-11", 55.02, 35.10, "RESMI BI", "26/275/DKom", "2024-12-18", "Siaran Pers RDG Desember 2024", BIEN.format("sp_2627524"), ""),
    ("2025-03", 56.30, 38.10, "PEJABAT BI", "", "2025-04-23", "Deputi Gubernur BI, konferensi pers RDG April 2025, via detik", "https://finance.detik.com/foto-bisnis/d-7889441/mantap-qris-tembus-56-juta-pengguna?zoom=1", "Rilis RDG 27/84/DKom hanya memuat pertumbuhan"),
    ("2025-06", 57.00, 39.30, "RESMI BI", "27/170/DKom", "2025-08-04", "Siaran Pers QRIS Jelajah Indonesia 2025", BI.format("sp_2717025"), "Posisi semester I 2025; 93,16% merchant UMKM"),
    ("2025-09", 58.00, 41.00, "PEJABAT BI", "", "2025-11-11", "Deputi Gubernur BI, Bulan Fintech Nasional, via Investortrust", "https://investortrust.id/financial/85155/tumbuh-158-transaksi-qris-tembus-10-33-miliar-per-september-2025", ""),
    ("2025-11", 59.00, 42.00, "PEJABAT BI", "", "2025-12-17", "Deputi Gubernur BI, konferensi pers RDG Desember 2025, via Bisnis.com", "https://ekonomi.bisnis.com/read/20251218/9/1937730/transaksi-qris-2-kali-lipat-dari-target-ekonomi-ri-cerah", "Disebut 'sampai dengan 2025'; posisi diasumsikan November 2025"),
    ("2026-01", 59.98, 43.06, "PEJABAT BI", "", "2026-03-12", "Deputi Gubernur BI via Kontan", "https://keuangan.kontan.co.id/news/bi-mencatat-nilai-transaksi-qris-capai-rp-16448-triliun-pada-januari-2026", "Pengguna +7,72% yoy, merchant +17,76% yoy"),
    ("2026-02", 60.77, None, "PEJABAT BI", "", "2026-04-08", "BI via Kontan", "https://keuangan.kontan.co.id/news/qris-cross-border-kian-diminati-transaksi-tumbuh-pesat", ""),
    ("2026-06", 65.77, 44.86, "RESMI BI", "28/159/DKom", "2026-08-17", "Siaran Pers Kartu Kredit Indonesia dan MDR 0%", BI.format("sp_2815926"), "6,23 juta pengguna baru Jan-Jun 2026; 96,68% merchant UMKM"),
    ("2026-07", None, 45.46, "RESMI BI", "28/191/DKom", "2026-09-22", "Siaran Pers QRIS Antarnegara 50-in-5 Awards", BI.format("sp_2819126"), "96,69% merchant UMKM"),
    ("2026-08", 69.32, 47.11, "RESMI BI", "28/196/DKom", "2026-09-24", "Siaran Pers FEKDI x IFSE 2026", BI.format("sp_2819626"), "Pengguna +20,26% yoy, merchant +13,83% yoy (yoy dari pernyataan Gubernur BI)"),
]

# periode, ukuran, volume (miliar trx), nilai (Rp triliun), status, no. rilis, sumber, url, catatan
TRX = [
    ("2020", "setahun", 0.12411, 8.21, "ASPI", "", "ASPI via RRI 18/04/2025", "https://www.rri.co.id/keuangan/1459019/pengguna-qris-makin-banyak-transaksi-capai-rp659-triliun", ""),
    ("2021", "setahun", 0.37469, 27.63, "ASPI", "", "ASPI via RRI 18/04/2025", "https://www.rri.co.id/keuangan/1459019/pengguna-qris-makin-banyak-transaksi-capai-rp659-triliun", ""),
    ("2022", "setahun", 1.00319, 99.98, "ASPI", "", "ASPI via RRI 18/04/2025", "https://www.rri.co.id/keuangan/1459019/pengguna-qris-makin-banyak-transaksi-capai-rp659-triliun", "BI resmi (25/211/DKom) menyebut 1,03 miliar transaksi"),
    ("2023", "setahun", 2.14, 226.00, "ASPI", "", "ASPI via RRI 18/04/2025", "https://www.rri.co.id/keuangan/1459019/pengguna-qris-makin-banyak-transaksi-capai-rp659-triliun", "BI resmi (26/9/DKom) menyebut nominal Rp229,96 triliun"),
    ("2024", "setahun", 6.24, 659.93, "ASPI", "", "ASPI via RRI 18/04/2025", "https://www.rri.co.id/keuangan/1459019/pengguna-qris-makin-banyak-transaksi-capai-rp659-triliun", "BI resmi (27/8/DKom) hanya menyebut pertumbuhan volume 175,2%"),
    ("2025", "setahun", 15.51, 1420.66, "PEJABAT BI", "", "Deputi Gubernur BI via Kontan (Mar 2026)", "https://keuangan.kontan.co.id/news/bi-mencatat-nilai-transaksi-qris-capai-rp-16448-triliun-pada-januari-2026", "Volume +148,54% yoy"),
    ("2025-H1", "Jan-Jun", 6.05, 579.00, "RESMI BI", "27/170/DKom", "Siaran Pers BI 04/08/2025", BI.format("sp_2717025"), ""),
    ("2026-H1", "Jan-Jun", 12.55, 1120.00, "RESMI BI", "28/159/DKom", "Siaran Pers BI 17/08/2026", BI.format("sp_2815926"), "Nilai Rp1,12 kuadriliun, +93,92% yoy"),
]

# periode, pertumbuhan yoy (%), basis, status, no. rilis, url
GROWTH = [
    ("2023-07", 84.50, "nominal", "25/233/DKom", BI.format("sp_2523323")),
    ("2023-08", 89.64, "nominal", "25/259/DKom", BIEN.format("sp_2525923")),
    ("2023-Q3", 87.90, "nominal", "25/283/DKom", BI.format("sp_2528323")),
    ("2023-10", 186.08, "nominal", "25/317/DKom", BIEN.format("sp_2531723")),
    ("2023-11", 157.43, "nominal", "25/342/DKom", BI.format("sp_2534223")),
    ("2023", 130.01, "nominal", "26/9/DKom", BI.format("sp_260924")),
    ("2024-01", 149.46, "nominal", "26/33/DKom", BI.format("sp_263324")),
    ("2024-02", 161.51, "nominal", "26/54/DKom", BI.format("sp_265424")),
    ("2024-Q1", 175.44, "nominal", "26/80/DKom", BI.format("sp_268024")),
    ("2024-04", 194.06, "nominal", "26/108/DKom", BI.format("sp_2610824")),
    ("2024-05", 213.31, "nominal", "26/125/DKom", BIEN.format("sp_2612524")),
    ("2024-Q2", 226.54, "tidak disebut", "26/144/DKom", BIEN.format("sp_2614424")),
    ("2024-08", 217.33, "tidak disebut", "26/201/DKom", BI.format("sp_2620124")),
    ("2024-Q3", 209.61, "tidak disebut", "26/225/DKom", BI.format("sp_2622524")),
    ("2024-10", 183.90, "tidak disebut", "26/254/DKom", BI.format("sp_2625424")),
    ("2024-11", 186.00, "volume", "26/275/DKom", BIEN.format("sp_2627524")),
    ("2024", 175.20, "volume", "27/8/DKom", BI.format("sp_270825")),
    ("2025-Q1", 169.15, "volume", "27/84/DKom", BI.format("sp_278425")),
    ("2025-05", 151.70, "volume", "27/133/DKom", BI.format("sp_2713325")),
    ("2025-07", 162.77, "volume", "27/194/DKom", BI.format("sp_2719425")),
    ("2025-08", 145.07, "volume", "27/219/DKom", BI.format("sp_2721925")),
    ("2025-10", 139.45, "tidak disebut", "27/274/DKom", BI.format("sp_2727425")),
    ("2026-05", 95.10, "volume", "28/126/DKom", BI.format("sp_2812626")),
    ("2026-Q2", 100.12, "tidak disebut", "28/142/DKom", BIEN.format("sp_2814226")),
    ("2026-08", 67.22, "volume", "28/193/DKom", BI.format("sp_2819326")),
]

# koridor, mulai, transaksi, nilai (Rp miliar), posisi, status, no. rilis, sumber, url
X = [
    ("Thailand", "2022-08", 994890, 437.54, "2025-06", "RESMI BI", "27/193/DKom", "Siaran Pers QRIS di Jepang", BI.format("sp_2719325")),
    ("Malaysia", "2023-05", 4310000, 1150.00, "2025-06", "RESMI BI", "27/193/DKom", "Siaran Pers QRIS di Jepang", BI.format("sp_2719325")),
    ("Singapura", "2023-11", 238216, 77.06, "2025-06", "RESMI BI", "27/193/DKom", "Siaran Pers QRIS di Jepang", BI.format("sp_2719325")),
    ("Thailand", "2022-08", 1640000, 656.27, "2026-02", "PEJABAT BI", "", "BI via Kontan 08/04/2026", "https://keuangan.kontan.co.id/news/qris-cross-border-kian-diminati-transaksi-tumbuh-pesat"),
    ("Singapura", "2023-11", 554510, 179.28, "2026-02", "PEJABAT BI", "", "BI via Kontan 08/04/2026", "https://keuangan.kontan.co.id/news/qris-cross-border-kian-diminati-transaksi-tumbuh-pesat"),
    ("Jepang (outbound)", "2025-08", 5088, 0.4288, "2026-02", "PEJABAT BI", "", "BI via Kontan 08/04/2026", "https://keuangan.kontan.co.id/news/qris-cross-border-kian-diminati-transaksi-tumbuh-pesat"),
    ("Total inbound", "2022-08", 21000000, 5900.00, "2026-07", "RESMI BI", "28/191/DKom", "Siaran Pers QRIS Antarnegara", BI.format("sp_2819126")),
    ("Total outbound", "2022-08", 4670000, 1730.00, "2026-07", "RESMI BI", "28/191/DKom", "Siaran Pers QRIS Antarnegara", BI.format("sp_2819126")),
]

HEAD = Font(name="Arial", bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F4E79")
BODY = Font(name="Arial", size=10)
LINK = Font(name="Arial", size=10, color="0563C1", underline="single")
STATUS_FILL = {"RESMI BI": "E2EFDA", "PEJABAT BI": "FFF2CC", "ASPI": "DDEBF7", "TURUNAN": "EDEDED"}

def sheet(ws, headers, rows, url_col=None, status_col=None, widths=None):
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, FILL, Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY
        if url_col is not None and row[url_col].value:
            row[url_col].hyperlink = row[url_col].value; row[url_col].font = LINK
        if status_col is not None:
            s = row[status_col].value
            if s in STATUS_FILL:
                row[status_col].fill = PatternFill("solid", fgColor=STATUS_FILL[s])
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"

wb = Workbook()
ws = wb.active; ws.title = "Pengguna_Merchant"
sheet(ws, ["Periode (posisi)", "Pengguna (juta)", "Merchant (juta)", "Pengguna per merchant", "Status sumber",
           "No. siaran pers", "Tanggal rilis", "Sumber", "URL", "Catatan"],
      [(p, u, m, None, s, n, d, src, url, note) for p, u, m, s, n, d, src, url, note in PM],
      url_col=8, status_col=4, widths=[14, 14, 14, 14, 13, 14, 12, 46, 50, 50])
for i in range(2, len(PM) + 2):
    ws[f"D{i}"] = f'=IF(AND(ISNUMBER(B{i}),ISNUMBER(C{i})),B{i}/C{i},"")'
    ws[f"D{i}"].number_format = "0.00"
    for col in "BC":
        ws[f"{col}{i}"].number_format = "0.00"

ws2 = wb.create_sheet("Transaksi")
sheet(ws2, ["Periode", "Cakupan", "Volume (miliar transaksi)", "Nilai (Rp triliun)", "Nilai rata-rata per transaksi (Rp ribu)",
            "Status sumber", "No. siaran pers", "Sumber", "URL", "Catatan"],
      [(p, c, v, n, None, s, no, src, url, note) for p, c, v, n, s, no, src, url, note in TRX],
      url_col=8, status_col=5, widths=[11, 10, 14, 14, 16, 13, 14, 38, 50, 50])
for i in range(2, len(TRX) + 2):
    ws2[f"E{i}"] = f"=IF(C{i}>0,D{i}/C{i},\"\")"
    ws2[f"E{i}"].number_format = "0.0"

ws3 = wb.create_sheet("Pertumbuhan_RDG")
sheet(ws3, ["Periode", "Pertumbuhan transaksi QRIS (% yoy)", "Basis", "Status sumber", "No. siaran pers", "URL"],
      [(p, g, b, "RESMI BI", n, url) for p, g, b, n, url in GROWTH],
      url_col=5, status_col=3, widths=[11, 16, 14, 13, 14, 60])

ws4 = wb.create_sheet("Antarnegara")
sheet(ws4, ["Koridor", "Mulai", "Transaksi kumulatif", "Nilai kumulatif (Rp miliar)", "Posisi data", "Status sumber",
            "No. siaran pers", "Sumber", "URL"], X, url_col=8, status_col=5,
      widths=[18, 10, 16, 16, 11, 13, 14, 32, 60])
ws4.append([]); r = ws4.max_row + 1
ws4[f"A{r}"] = "Total antarnegara (inbound + outbound, s.d. Jul 2026)"
ws4[f"C{r}"] = "=C8+C9"; ws4[f"D{r}"] = "=D8+D9"
ws4[f"A{r+1}"] = "Porsi inbound (transaksi)"; ws4[f"C{r+1}"] = "=C8/(C8+C9)"; ws4[f"C{r+1}"].number_format = "0.0%"
ws4[f"A{r+2}"] = "Surplus nilai inbound - outbound (Rp miliar)"; ws4[f"D{r+2}"] = "=D8-D9"
for rr in range(r, r + 3):
    for c in ws4[rr]:
        c.font = Font(name="Arial", size=10, bold=True)

ws5 = wb.create_sheet("Keterangan")
info = [
    ["Dataset QRIS Indonesia 2021–2026"],
    ["Penyusun: Ade Rumintarsih, Boash University, Bogor. Disusun 12 Oktober 2026."],
    [""],
    ["Status sumber", "Arti"],
    ["RESMI BI", "Angka tertulis langsung di siaran pers Bank Indonesia (bi.go.id). Paling kuat untuk dikutip."],
    ["PEJABAT BI", "Pernyataan Gubernur/Deputi Gubernur/pejabat BI di konferensi pers atau acara, dikutip media. Kutip media dan sebut pejabatnya."],
    ["ASPI", "Data Asosiasi Sistem Pembayaran Indonesia yang dikutip media."],
    ["TURUNAN", "Dihitung dari angka resmi (cara hitung ada di kolom Catatan)."],
    [""],
    ["Catatan penting"],
    ["1. Tabel Statistik Sistem Pembayaran dan Infrastruktur Pasar Keuangan (SPIP) BI tidak memuat baris QRIS; sumber resmi utama adalah siaran pers BI."],
    ["2. Sejak 2025, siaran pers RDG hanya memuat pertumbuhan transaksi QRIS, bukan jumlah pengguna/merchant; angka 2025–2026 diambil dari siaran pers tematik BI atau pernyataan pejabat BI."],
    ["3. Siaran pers RDG tidak selalu menyebut bulan posisi; posisi ditetapkan dari konteks paragraf (lihat kolom Catatan)."],
    ["4. Angka pengguna adalah pengguna terdaftar, bukan pengguna aktif."],
    ["5. Kolom D (Pengguna_Merchant) dan E (Transaksi) serta ringkasan Antarnegara adalah rumus; ubah angka di kolom input dan hasilnya ikut berubah."],
]
for row in info:
    ws5.append(row)
ws5["A1"].font = Font(name="Arial", bold=True, size=13)
for row in ws5.iter_rows(min_row=2):
    for c in row:
        if c.font.size != 13:
            c.font = BODY
for row in (4, 10):
    ws5[f"A{row}"].font = Font(name="Arial", bold=True)
for s, f in STATUS_FILL.items():
    for rr in range(5, 9):
        if ws5[f"A{rr}"].value == s:
            ws5[f"A{rr}"].fill = PatternFill("solid", fgColor=f)
ws5.column_dimensions["A"].width = 22; ws5.column_dimensions["B"].width = 110

wb.move_sheet("Keterangan", offset=-4)
wb.active = 0
wb.save(f"{OUT}/Dataset_QRIS_Indonesia_2021-2026.xlsx")

with open(f"{OUT}/qris_pengguna_merchant.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["periode", "pengguna_juta", "merchant_juta", "status", "sumber"])
    for p, u, m, s, n, d, src, url, note in PM:
        w.writerow([p, "" if u is None else u, "" if m is None else m, s, f"{src} {n} {url}".strip()])
with open(f"{OUT}/qris_transaksi_tahunan.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["periode", "cakupan", "volume_miliar_trx", "nilai_triliun_rp", "status", "sumber"])
    for p, c, v, n, s, no, src, url, note in TRX:
        w.writerow([p, c, v, n, s, f"{src} {no} {url}".strip()])
with open(f"{OUT}/qris_antarnegara.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["koridor", "mulai", "transaksi", "nilai_miliar_rp", "posisi", "status", "sumber"])
    for k, m, t, v, pos, s, no, src, url in X:
        w.writerow([k, m, t, v, pos, s, f"{src} {no} {url}".strip()])
with open(f"{OUT}/qris_pertumbuhan_rdg.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["periode", "pertumbuhan_yoy_persen", "basis", "no_siaran_pers", "url"])
    for row in GROWTH:
        w.writerow(row)
print("pengguna:", sum(1 for r in PM if r[1] is not None), "merchant:", sum(1 for r in PM if r[2] is not None),
      "resmi pengguna:", sum(1 for r in PM if r[1] is not None and r[3] == "RESMI BI"),
      "resmi merchant:", sum(1 for r in PM if r[2] is not None and r[3] == "RESMI BI"))
